from __future__ import annotations

import io
import json

import pandas as pd
from audit_api.deps import get_pipeline, get_store
from audit_api.main import app
from fastapi.testclient import TestClient
from openpyxl import load_workbook

client = TestClient(app)


def _minimal_xlsx() -> io.BytesIO:
    df = pd.DataFrame(
        {
            "凭证编号": ["1001", "1001", "2001", "2001"],
            "过账日期": ["2023-12-28", "2023-12-28", "2024-03-15", "2024-03-15"],
            "凭证货币价值": [500000.0, -500000.0, 1000.0, -1000.0],
            "借/贷标识": ["S", "H", "S", "H"],
            "总账科目": ["6001010000", "1122010000", "6001010000", "6401010000"],
            "总账科目：短文本": ["主营业务收入", "应收账款", "主营业务收入", "主营业务成本"],
            "凭证类型": ["SA", "SA", "AF", "AF"],
            "文本": ["销售", "销售", "成本", "成本"],
        }
    )
    buf = io.BytesIO()
    df.to_excel(buf, index=False, engine="openpyxl")
    buf.seek(0)
    buf.name = "test.xlsx"  # type: ignore[attr-defined]
    return buf


def _ingest_two_years(tmp_path, monkeypatch) -> str:
    monkeypatch.setenv("AUDIT_WORKBENCH_DATA_ROOT", str(tmp_path))
    get_store.cache_clear()
    get_pipeline.cache_clear()

    pid = client.post("/projects", json={"name": "pipeline测试"}).json()["project_id"]

    files = []
    for year, month in [(2023, "12-28"), (2024, "03-15")]:
        df = pd.DataFrame(
            {
                "凭证编号": [f"{year}001", f"{year}001"],
                "过账日期": [f"{year}-{month}", f"{year}-{month}"],
                "凭证货币价值": [500000.0, -500000.0],
                "借/贷标识": ["S", "H"],
                "总账科目": ["6001010000", "1122010000"],
                "总账科目：短文本": ["主营业务收入", "应收账款"],
                "凭证类型": ["SA", "SA"],
                "文本": ["销售", "销售"],
            }
        )
        buf = io.BytesIO()
        df.to_excel(buf, index=False, engine="openpyxl")
        buf.seek(0)
        files.append((f"{year}.xlsx", buf.getvalue()))

    det = client.post(
        f"/projects/{pid}/ingest/detect",
        files=[("files", (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for name, data in files],
    ).json()

    commit = client.post(
        f"/projects/{pid}/ingest/commit",
        files=[("files", (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for name, data in files],
        data={"column_mapping": json.dumps(det["suggested_mapping"])},
    )
    assert commit.status_code == 200
    assert len(commit.json()["years"]) >= 2

    get_store.cache_clear()
    get_pipeline.cache_clear()
    return pid


def test_pipeline_profiles_rules_samples_export(tmp_path, monkeypatch) -> None:
    pid = _ingest_two_years(tmp_path, monkeypatch)

    prof = client.post(f"/projects/{pid}/pipeline/profiles")
    assert prof.status_code == 200
    assert prof.json()["years"]

    rules = client.post(f"/projects/{pid}/pipeline/rules/run")
    assert rules.status_code == 200
    rule_context = rules.json()["rule_run_context"]
    assert rule_context["rule_run_id"].startswith("rr_")
    assert len(rule_context["result_hash"]) == 64
    assert rule_context["ingest_run_id"].startswith("ing_")
    assert rule_context["engine_revision"]
    assert rule_context["rules_snapshot"]
    assert rule_context["result_artifact"].endswith(".json.gz")
    store = get_store()
    state = store.load_state(pid)
    assert all(not block["hits"] for block in state["rule_results"])
    artifact = store.project_dir(pid) / rule_context["result_artifact"]
    assert artifact.exists()
    full_results = store.load_rule_run_results(pid, state=state)
    assert sum(len(block["hits"]) for block in full_results) == rule_context["complete_hit_count"]
    assert all(len(block["hits"]) <= 20 for block in rules.json()["results"])

    cross = client.post(f"/projects/{pid}/pipeline/cross-year")
    assert cross.status_code == 200

    samples = client.post(
        f"/projects/{pid}/pipeline/samples",
        json={"method": "by_rule", "size": 10},
    )
    assert samples.status_code == 200

    export = client.get(f"/projects/{pid}/pipeline/export")
    assert export.status_code == 200
    assert "filename*=UTF-8''" in export.headers["content-disposition"]
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert len(export.content) > 1000

    # 导出必须服从页面当前样本事实源，而不是重新按规则生成另一套样本。
    store = get_store()
    state = store.load_state(pid)
    selected_voucher = "2023001"
    state["samples"] = [{"凭证编号": selected_voucher}]
    store.save_state(pid, state)
    export = client.get(f"/projects/{pid}/pipeline/export")
    workbook = load_workbook(io.BytesIO(export.content), read_only=True)
    sample_sheet = workbook["样本清单"]
    exported_vouchers = {
        str(sample_sheet.cell(row=row, column=2).value)
        for row in range(2, sample_sheet.max_row + 1)
        if sample_sheet.cell(row=row, column=2).value
    }
    assert exported_vouchers == {selected_voucher}

    get_store.cache_clear()
    get_pipeline.cache_clear()


def test_versioned_full_population_random_sample_is_replayable(tmp_path, monkeypatch) -> None:
    pid = _ingest_two_years(tmp_path, monkeypatch)
    body = {
        "plan_name": "完整总体随机测试",
        "population_scope": "full_population",
        "strategy": "random",
        "method": "random",
        "size": 2,
        "seed": 20260730,
        "years": [2023, 2024],
        "coverage_constraints": {},
    }

    first = client.post(f"/projects/{pid}/pipeline/samples", json=body)
    second = client.post(f"/projects/{pid}/pipeline/samples", json=body)

    assert first.status_code == 200
    assert second.status_code == 200
    payload = first.json()
    assert payload["requested_plan"]["population_scope"] == "full_population"
    assert payload["population_snapshot"]["voucher_count"] == 2
    assert payload["population_snapshot"]["population_membership_revision"]
    assert payload["population_snapshot"]["ingest_run_id"].startswith("ing_")
    assert payload["population_snapshot"]["engine_revision"]
    assert payload["selection_trace"]["selection_id"] == second.json()["selection_trace"]["selection_id"]
    assert payload["selection_trace"]["inclusion_probability_available"] is True
    assert payload["selection_trace"]["statistical_projection_allowed"] is True
    assert payload["selection_trace"]["uniform_inclusion_probability"] == 1.0
    assert payload["selection_trace"]["selected_voucher_keys"] == [
        "VK|__NO_COMPANY__|2023|2023001",
        "VK|__NO_COMPANY__|2024|2024001",
    ]
    assert len(payload["selection_trace"]["selected_voucher_keys_digest"]) == 64
    assert {row["_voucher_key"] for row in payload["samples"]} == {
        "VK|__NO_COMPANY__|2023|2023001",
        "VK|__NO_COMPANY__|2024|2024001",
    }

    cached = client.get(f"/projects/{pid}/pipeline/samples")
    assert cached.status_code == 200
    assert cached.json()["selection_trace"]["selection_id"] == payload["selection_trace"]["selection_id"]

    get_store.cache_clear()
    get_pipeline.cache_clear()


def test_put_rules_invalidates_derived_results(tmp_path, monkeypatch) -> None:
    pid = _ingest_two_years(tmp_path, monkeypatch)

    assert client.post(f"/projects/{pid}/pipeline/rules/run").status_code == 200
    assert client.post(f"/projects/{pid}/pipeline/cross-year").status_code == 200
    assert client.post(
        f"/projects/{pid}/pipeline/samples",
        json={"method": "by_rule", "size": 10},
    ).status_code == 200

    before = client.get(f"/projects/{pid}/rules").json()
    assert before["splitting"]["enabled"] is True
    assert client.get(f"/projects/{pid}/pipeline/rules/results").json()["total_hits"] >= 0
    assert client.get(f"/projects/{pid}/pipeline/cross-year").json()["count"] >= 0
    assert client.get(f"/projects/{pid}/pipeline/samples").json()["sample_rows"] >= 0

    before["splitting"] = {**before["splitting"], "enabled": False}
    before["cross_year_accrual"] = {
        **before["cross_year_accrual"],
        "coverage_threshold": 0.55,
    }
    put = client.put(f"/projects/{pid}/rules", json=before)
    assert put.status_code == 200
    assert put.json()["splitting"]["enabled"] is False

    store = get_store()
    state = store.load_state(pid)
    assert state.get("rule_results") == []
    assert "rule_run_context" not in state
    assert state.get("samples") == []
    assert "sampling_plan" not in state
    assert state.get("cross_year_findings") == []
    assert state.get("llm_judgments") == {}

    get_store.cache_clear()
    get_pipeline.cache_clear()


def test_sampling_rejects_tampered_rule_fact_artifact(tmp_path, monkeypatch) -> None:
    pid = _ingest_two_years(tmp_path, monkeypatch)
    assert client.post(f"/projects/{pid}/pipeline/rules/run").status_code == 200
    store = get_store()
    state = store.load_state(pid)
    artifact = store.project_dir(pid) / state["rule_run_context"]["result_artifact"]
    artifact.write_bytes(b"not-a-gzip-rule-result")

    response = client.post(
        f"/projects/{pid}/pipeline/samples",
        json={
            "population_scope": "risk_signals",
            "strategy": "risk_directed",
            "method": "by_rule",
            "size": 10,
        },
    )
    assert response.status_code == 400
    assert "规则结果事实文件损坏" in response.json()["detail"]

    get_store.cache_clear()
    get_pipeline.cache_clear()


def test_sampling_rejects_unsupported_or_unmet_explicit_constraints(
    tmp_path,
    monkeypatch,
) -> None:
    pid = _ingest_two_years(tmp_path, monkeypatch)

    unsupported = client.post(
        f"/projects/{pid}/pipeline/samples",
        json={
            "population_scope": "full_population",
            "strategy": "random",
            "method": "random",
            "size": 2,
            "coverage_constraints": {"max_same_risk_signal_ratio": 0.5},
        },
    )
    assert unsupported.status_code == 400
    assert "仅适用于风险定向抽样" in unsupported.json()["detail"]

    infeasible = client.post(
        f"/projects/{pid}/pipeline/samples",
        json={
            "population_scope": "full_population",
            "strategy": "random",
            "method": "random",
            "size": 2,
            "coverage_constraints": {"min_per_month": 2},
        },
    )
    assert infeasible.status_code == 400
    assert "覆盖约束不可满足" in infeasible.json()["detail"]
    assert "sampling_plan" not in get_store().load_state(pid)

    get_store.cache_clear()
    get_pipeline.cache_clear()
