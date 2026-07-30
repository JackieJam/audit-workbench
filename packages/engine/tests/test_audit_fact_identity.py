from __future__ import annotations

import io

import pandas as pd
import pytest
from audit_engine.candidate_pool import (
    _build_voucher_amount_map,
    sample_from_pool,
    sample_from_rule_results,
    samples_for_voucher_ids,
)
from audit_engine.data_columns import (
    LINE_KEY_COLUMN,
    VOUCHER_KEY_COLUMN,
    audit_input_quality,
    ensure_analysis_columns,
    ensure_voucher_identity,
)
from audit_engine.ingestion import load_files
from audit_engine.reporter import generate_report_bytes
from audit_engine.rule_engine import run_all_rules
from openpyxl import load_workbook


def _journal_rows() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for company, year, amount in (
        ("C01", 2023, 100.0),
        ("C01", 2024, 200.0),
        ("C02", 2023, 300.0),
        ("C02", 2024, 400.0),
    ):
        for dc, account in (("S", "100201"), ("H", "600101")):
            rows.append({
                "公司代码": company,
                "会计年度": year,
                "凭证编号": "0001",
                "过账日期": pd.Timestamp(year=year, month=3, day=15),
                "借/贷标识": dc,
                "凭证货币价值": amount,
                "凭证货币代码": "CNY",
                "总账科目": account,
                "总账科目：长文本": "银行存款" if dc == "S" else "主营业务收入",
                "凭证类型": "SA",
                "文本": "测试",
                "用户名": "tester",
                "供应商编号": "",
            })
    return pd.DataFrame(rows)


def test_voucher_and_line_identity_never_merge_same_number_across_years_or_companies() -> None:
    work = ensure_analysis_columns(_journal_rows())

    assert work[VOUCHER_KEY_COLUMN].nunique() == 4
    assert work[LINE_KEY_COLUMN].nunique() == 8
    assert all(
        {"C01", "C02"}.intersection(str(key).split("|"))
        for key in work[VOUCHER_KEY_COLUMN]
    )

    target = work.loc[
        work["公司代码"].eq("C01") & work["会计年度"].eq(2024),
        VOUCHER_KEY_COLUMN,
    ].iloc[0]
    samples = samples_for_voucher_ids({target}, work)
    assert len(samples) == 2
    assert {row["会计年度"] for row in samples} == {2024}
    assert {row["公司代码"] for row in samples} == {"C01"}
    assert {row["_voucher_key"] for row in samples} == {target}


def test_missing_year_is_explicitly_blocked_and_never_silently_grouped() -> None:
    frame = pd.DataFrame({
        "公司代码": ["C01", "C01"],
        "凭证编号": ["0001", "0001"],
        "借/贷标识": ["S", "H"],
        "凭证货币价值": [100.0, 100.0],
        "总账科目": ["100201", "600101"],
    })
    identified = ensure_voucher_identity(frame)
    assert identified[VOUCHER_KEY_COLUMN].nunique() == 2
    assert set(identified["_voucher_identity_quality"]) == {"missing_year"}

    quality = audit_input_quality(frame)
    assert quality["status"] == "blocked"
    assert "missing_year_source" in {
        issue["code"] for issue in quality["issues"]
    }


def test_normalized_voucher_amount_does_not_double_count_debit_and_credit() -> None:
    work = ensure_analysis_columns(_journal_rows().iloc[:2].copy())
    key = str(work[VOUCHER_KEY_COLUMN].iloc[0])
    amount_map = _build_voucher_amount_map(work)
    assert amount_map[key] == 100.0


def test_rule_hits_are_complete_and_sample_size_only_limits_selection() -> None:
    frame = pd.DataFrame({
        "公司代码": ["C01"] * 75,
        "会计年度": [2024] * 75,
        "凭证编号": [f"V{i:03d}" for i in range(75)],
        "过账日期": pd.to_datetime(["2024-03-15"] * 75),
        "借/贷标识": ["S"] * 75,
        "凭证货币价值": [1_000_000.0] * 75,
        "凭证货币代码": ["CNY"] * 75,
        "总账科目": ["660201"] * 75,
        "总账科目：长文本": ["管理费用"] * 75,
        "凭证类型": ["SA"] * 75,
        "文本": ["大额测试"] * 75,
        "用户名": [f"u{i:03d}" for i in range(75)],
        "供应商编号": [""] * 75,
    })
    cfg = {
        "max_sample_size": 10,
        "large_amount": {
            "enabled": True,
            "round_number_threshold": 1_000_000,
            "repeat_threshold": 10**20,
            "holiday_min_amount": 10**20,
        },
    }
    results = run_all_rules(frame, cfg)

    assert sum(result.count for result in results) == 75
    samples = sample_from_rule_results(results, frame, size=10)
    assert len({row["_voucher_key"] for row in samples}) == 10
    assert all(row["风险评分"] is not None for row in samples)
    assert all(row["入样理由"] for row in samples)
    assert all(row["评分构成"] for row in samples)


def test_individual_rules_do_not_hide_hits_behind_legacy_display_caps() -> None:
    sensitive = pd.DataFrame({
        "公司代码": ["C01"] * 250,
        "会计年度": [2024] * 250,
        "凭证编号": [f"S{i:03d}" for i in range(250)],
        "过账日期": pd.to_datetime(["2024-03-15"] * 250),
        "借/贷标识": ["S"] * 250,
        "凭证货币价值": [20_000.0] * 250,
        "凭证货币代码": ["CNY"] * 250,
        "总账科目": ["660201"] * 250,
        "总账科目：长文本": ["管理费用-咨询费"] * 250,
        "凭证类型": ["SA"] * 250,
        "文本": ["咨询服务"] * 250,
        "用户名": [f"u{i:03d}" for i in range(250)],
        "供应商编号": [""] * 250,
    })
    sensitive_results = run_all_rules(
        sensitive,
        {
            "sensitive_fees": {
                "enabled": True,
                "categories": {
                    "咨询费": {
                        "keywords": ["咨询"],
                        "exclude": [],
                        "threshold": 10_000,
                    }
                },
            }
        },
    )
    assert sum(result.count for result in sensitive_results) == 250

    trade_rows: list[dict[str, object]] = []
    for index in range(60):
        voucher_id = f"T{index:03d}"
        trade_rows.extend([
            {
                "公司代码": "C01",
                "会计年度": 2024,
                "凭证编号": voucher_id,
                "过账日期": pd.Timestamp("2024-06-15"),
                "借/贷标识": "H",
                "凭证货币价值": 1_000_000.0,
                "凭证货币代码": "CNY",
                "总账科目": "600101",
                "总账科目：长文本": "主营业务收入",
                "凭证类型": "SA",
                "文本": "设备销售",
                "用户名": "sales",
                "供应商编号": "",
            },
            {
                "公司代码": "C01",
                "会计年度": 2024,
                "凭证编号": voucher_id,
                "过账日期": pd.Timestamp("2024-06-15"),
                "借/贷标识": "S",
                "凭证货币价值": 990_000.0,
                "凭证货币代码": "CNY",
                "总账科目": "640101",
                "总账科目：长文本": "主营业务成本",
                "凭证类型": "SA",
                "文本": "设备销售",
                "用户名": "sales",
                "供应商编号": "",
            },
        ])
    trade_results = run_all_rules(
        pd.DataFrame(trade_rows),
        {
            "financing_trade": {
                "enabled": True,
                "min_revenue_amount": 1,
                "low_margin_threshold": 0.05,
                "max_candidate_groups": 50,
            }
        },
    )
    assert sum(result.count for result in trade_results) == 60


def test_amount_weighted_sampling_blocks_mixed_document_currencies() -> None:
    frame = _journal_rows().iloc[:4].copy()
    frame.loc[frame.index[:2], "凭证货币代码"] = "CNY"
    frame.loc[frame.index[2:], "凭证货币代码"] = "USD"
    pool = [{
        "status": "候选",
        "voucher_keys": ensure_voucher_identity(frame)[VOUCHER_KEY_COLUMN]
        .drop_duplicates()
        .tolist(),
        "voucher_ids": ["0001"],
    }]

    with pytest.raises(ValueError, match="多种凭证币"):
        run_all_rules(
            frame,
            {"large_amount": {"enabled": True}},
        )
    with pytest.raises(ValueError, match="多种凭证币"):
        sample_from_pool(pool, frame, method="monetary_unit", size=1)


def test_ingestion_coordinates_flow_into_stable_line_key() -> None:
    source = _journal_rows().iloc[:2].drop(
        columns=["会计年度"],
    )
    buffer = io.BytesIO()
    source.to_excel(buffer, index=False, engine="openpyxl", sheet_name="序时账")
    buffer.seek(0)
    buffer.name = "journal.xlsx"  # type: ignore[attr-defined]
    buffer._audit_source_asset = {  # type: ignore[attr-defined]
        "asset_id": "src_test",
        "sha256": "abc123",
        "original_name": "journal.xlsx",
    }

    unified, _, _ = load_files([buffer])
    work = ensure_voucher_identity(unified)
    assert set(work["_source_asset_id"]) == {"src_test"}
    assert set(work["_source_file_hash"]) == {"abc123"}
    assert set(work["_source_sheet"]) == {"序时账"}
    assert work["_source_row"].tolist() == [2, 3]
    assert work[LINE_KEY_COLUMN].nunique() == 2


def test_export_uses_explicit_stable_key_without_cross_year_expansion() -> None:
    frame = _journal_rows()
    work = ensure_voucher_identity(frame)
    selected_key = str(
        work.loc[
            work["公司代码"].eq("C02") & work["会计年度"].eq(2024),
            VOUCHER_KEY_COLUMN,
        ].iloc[0]
    )
    data, stats = generate_report_bytes(
        frame,
        [],
        explicit_samples=[{"_voucher_key": selected_key, "凭证编号": "0001"}],
        max_sample_size=10,
    )

    workbook = load_workbook(io.BytesIO(data), read_only=True)
    rows = list(workbook["样本清单"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    assert stats["sample_vouchers"] == 1
    assert len(rows) == 2
    assert {row[4] for row in rows} == {2024}
