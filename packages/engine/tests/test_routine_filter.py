"""常规机械分录过滤与导出压缩测试。"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from audit_engine.cross_year import _counterparty_circular_flow
from audit_engine.reporter import generate_report_bytes
from audit_engine.routine_filter import (
    apply_routine_exclusion,
    prepare_candidate_detail,
    prioritize_voucher_ids_for_export,
    pure_routine_voucher_ids,
)
from audit_engine.rule_engine import RuleHit, RuleResult, apply_whitelist
from audit_engine.rules_config import default_rules_config
from openpyxl import load_workbook


def _cfg() -> dict:
    return default_rules_config()


def test_whitelist_matches_account_name_not_only_text() -> None:
    df = pd.DataFrame(
        {
            "凭证编号": ["A1", "A2"],
            "凭证类型": ["SA", "SA"],
            "凭证货币价值": [10000.0, 20000.0],
            "文本": ["202401计提人工费", "支付咨询费"],
            "总账科目：长文本": ["费用-人工-社会保险费-基本养老保险", "管理费用-咨询费"],
            "总账科目": ["6910010701", "6602010001"],
        }
    )
    kept, excluded = apply_whitelist(df, _cfg())
    assert "A1" in set(excluded["凭证编号"].astype(str))
    assert "A2" in set(kept["凭证编号"].astype(str))


def test_prepare_candidate_detail_drops_statutory_and_caps() -> None:
    rows = []
    for i in range(60):
        rows.append(
            {
                "凭证编号": f"V{i:03d}",
                "凭证类型": "SA",
                "凭证货币价值": 1_000_000 - i * 1000,
                "_amount_abs": 1_000_000 - i * 1000,
                "文本": "市场调研",
                "总账科目：长文本": "管理费用-市场调研费",
                "总账科目": "6602090001",
            }
        )
    # statutory noise vouchers
    for i in range(5):
        rows.append(
            {
                "凭证编号": f"S{i}",
                "凭证类型": "SA",
                "凭证货币价值": 50000,
                "_amount_abs": 50000,
                "文本": "202401计提人工费",
                "总账科目：长文本": "费用-人工-社会保险费-医疗保险",
                "总账科目": ["6910010702"][0],
            }
        )
    detail = pd.DataFrame(rows)
    cleaned = prepare_candidate_detail(detail, _cfg(), selector={"expense_category": "市场调研"})
    vids = set(cleaned["凭证编号"].astype(str))
    assert not any(v.startswith("S") for v in vids)
    assert len(vids) <= 50


def test_selector_protects_fx_category() -> None:
    df = pd.DataFrame(
        {
            "凭证编号": ["FX1"],
            "凭证类型": "SA",
            "凭证货币价值": [100000.0],
            "_amount_abs": [100000.0],
            "文本": ["月末调汇"],
            "总账科目：长文本": ["财务费用-汇兑损益-调汇"],
            "总账科目": ["6603049901"],
        }
    )
    kept, _ = apply_routine_exclusion(
        df, _cfg(), purpose="candidates", selector={"expense_category": "财务费用(汇兑)"}
    )
    assert len(kept) == 1


def test_export_respects_max_sample_size_and_skips_routine_lines() -> None:
    rows = []
    for i in range(80):
        rows.append(
            {
                "凭证编号": f"H{i:03d}",
                "过账日期": "2024-01-15",
                "凭证类型": "SA",
                "凭证货币价值": 100000 + i,
                "借/贷标识": "S",
                "总账科目": "6602010001",
                "总账科目：长文本": "管理费用-咨询费",
                "文本": "咨询费",
                "供应商科目：名称 1": "",
                "客户科目：姓名 1": "",
                "用户名": "u1",
                "_year": 2024,
                "_amount_raw": 100000 + i,
            }
        )
        # sibling statutory line on same voucher
        rows.append(
            {
                "凭证编号": f"H{i:03d}",
                "过账日期": "2024-01-15",
                "凭证类型": "SA",
                "凭证货币价值": 8000,
                "借/贷标识": "S",
                "总账科目": "6910010701",
                "总账科目：长文本": "费用-人工-社会保险费-基本养老保险",
                "文本": "202401计提人工费",
                "供应商科目：名称 1": "",
                "客户科目：姓名 1": "",
                "用户名": "u1",
                "_year": 2024,
                "_amount_raw": 8000,
            }
        )
    df = pd.DataFrame(rows)
    explicit = [{"凭证编号": f"H{i:03d}"} for i in range(80)]
    hits = [
        RuleResult(
            rule_name="敏感费用筛查",
            hits=[RuleHit(voucher_id=f"H{i:03d}", rule_type="敏感费用(咨询费)", evidence="x", line_indices=(), priority=3)],
        )
        for i in range(10)
    ]
    data, stats = generate_report_bytes(
        df,
        hits,
        max_sample_size=20,
        explicit_samples=explicit,
        rules_config=_cfg(),
    )
    assert stats["sample_vouchers"] == 20
    assert len(data) > 1000
    # exported workbook should not be dominated by social insurance lines
    wb = load_workbook(BytesIO(data), read_only=True)
    ws = wb["样本清单"]
    names = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        names.append(str(row[8] or ""))
    wb.close()
    assert sum("社会保险" in n for n in names) == 0
    assert sum("咨询费" in n for n in names) > 0


def test_prioritize_hit_vouchers() -> None:
    ordered = prioritize_voucher_ids_for_export(
        ["A", "B", "C", "D"],
        hit_voucher_ids={"C", "A"},
        max_size=3,
    )
    assert ordered == ["A", "C", "B"]


def test_circular_skips_empty_vendor_and_caps() -> None:
    year_map = {
        2023: pd.DataFrame(
            {
                "过账日期": pd.to_datetime(["2023-12-20"] * 30),
                "凭证货币价值": [-1_000_000.0] * 30,
                "供应商编号": [""] * 30,
                "凭证编号": [f"O{i}" for i in range(30)],
                "供应商科目：名称 1": [""] * 30,
            }
        ),
        2024: pd.DataFrame(
            {
                "过账日期": pd.to_datetime(["2024-01-10"] * 30),
                "凭证货币价值": [900_000.0] * 30,
                "供应商编号": [""] * 30,
                "凭证编号": [f"I{i}" for i in range(30)],
                "供应商科目：名称 1": [""] * 30,
            }
        ),
    }
    assert _counterparty_circular_flow(year_map) == []

    year_map[2023]["供应商编号"] = ["V1"] * 30
    year_map[2024]["供应商编号"] = ["V1"] * 30
    year_map[2023]["供应商科目：名称 1"] = ["甲公司"] * 30
    findings = _counterparty_circular_flow(year_map, max_vouchers=5)
    assert len(findings) == 1
    assert len(findings[0].voucher_ids) == 5
    assert findings[0].evidence.get("vendor") == "V1"


def test_pure_routine_voucher_detection() -> None:
    df = pd.DataFrame(
        {
            "凭证编号": ["P1", "P1", "G1", "G1"],
            "凭证类型": ["SA", "SA", "SA", "SA"],
            "凭证货币价值": [10000, -10000, 50000, -50000],
            "文本": ["社保缴纳", "社保缴纳", "支付咨询费", "支付咨询费"],
            "总账科目：长文本": [
                "应付职工薪酬-社会保险费",
                "银行存款-活期-人民币",
                "管理费用-咨询费",
                "银行存款-活期-人民币",
            ],
            "总账科目": ["2211030001", "1002010001", "6602010001", "1002010001"],
        }
    )
    pure = pure_routine_voucher_ids(df, _cfg(), purpose="candidates", threshold=0.5)
    assert "P1" in pure
    assert "G1" not in pure
